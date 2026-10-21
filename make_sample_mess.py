"""Generates a realistic 'sample_mess/' directory with ~25 messy files

including Pillow-rendered screenshot images, docx files, PDFs, receipts,
code scripts, and notes for demoing Sloth Sorter.
"""

import io
import os
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx


def create_pdf(filepath: Path, title: str, paragraphs: list[str]):
    """Creates a clean, valid standard PDF with text stream readable by pypdf."""
    stream_lines = [
        "BT",
        "/F1 16 Tf",
        f"50 720 Td ({title}) Tj",
        "/F1 11 Tf",
        "0 -30 Td",
    ]
    for p in paragraphs:
        # Escape parens for pdf string
        safe_p = p.replace("(", "\\(").replace(")", "\\)")
        stream_lines.append(f"({safe_p}) Tj")
        stream_lines.append("0 -20 Td")
    stream_lines.append("ET")

    stream_content = "\n".join(stream_lines)
    stream_bytes = stream_content.encode("ascii", errors="replace")

    # Construct minimal valid PDF structure
    obj1 = b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
    obj2 = b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
    obj3 = b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n"
    obj4 = b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
    obj5 = f"5 0 obj << /Length {len(stream_bytes)} >> stream\n".encode("ascii") + stream_bytes + b"\nendstream\nendobj\n"

    header = b"%PDF-1.4\n"
    pos1 = len(header)
    pos2 = pos1 + len(obj1)
    pos3 = pos2 + len(obj2)
    pos4 = pos3 + len(obj3)
    pos5 = pos4 + len(obj4)
    xref_pos = pos5 + len(obj5)

    xref = f"""xref
0 6
0000000000 65535 f 
{pos1:010d} 00000 n 
{pos2:010d} 00000 n 
{pos3:010d} 00000 n 
{pos4:010d} 00000 n 
{pos5:010d} 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
{xref_pos}
%%EOF"""

    with open(filepath, "wb") as f:
        f.write(header + obj1 + obj2 + obj3 + obj4 + obj5 + xref.encode("ascii"))


def create_image_with_ui(filepath: Path, width=700, height=450, bg="#1E1E2E", title="Slack - Bug Discussion", lines=None):
    """Generates a realistic screenshot graphic with window chrome and text."""
    img = Image.new("RGB", (width, height), color=bg)
    draw = ImageDraw.Draw(img)

    # Window titlebar
    draw.rectangle([(0, 0), (width, 36)], fill="#181825")
    # Window traffic lights
    draw.ellipse([(14, 12), (24, 22)], fill="#FF5F56")
    draw.ellipse([(32, 12), (42, 22)], fill="#FFBD2E")
    draw.ellipse([(50, 12), (60, 22)], fill="#27C93F")

    # Title text
    draw.text((80, 10), title, fill="#CDD6F4")

    # Draw body lines
    y = 60
    if lines:
        for line, color in lines:
            draw.text((30, y), line, fill=color)
            y += 24

    img.save(filepath)


def create_nature_photo(filepath: Path, width=640, height=480):
    """Draws a scenic forest / sloth landscape photo."""
    img = Image.new("RGB", (width, height), color="#2b4a33")
    draw = ImageDraw.Draw(img)

    # Sky gradient / top
    draw.rectangle([(0, 0), (width, int(height * 0.4))], fill="#87CEEB")
    # Sun
    draw.ellipse([(width - 120, 30), (width - 40, 110)], fill="#FFA500")
    # Mountain shapes
    draw.polygon([(0, int(height * 0.45)), (200, 100), (400, int(height * 0.45))], fill="#4A6B53")
    draw.polygon([(250, int(height * 0.45)), (480, 120), (width, int(height * 0.45))], fill="#3D5A44")
    # Forest floor
    draw.rectangle([(0, int(height * 0.45)), (width, height)], fill="#1B3B22")
    # Tree trunks
    for x in range(40, width, 80):
        draw.rectangle([(x, int(height * 0.3)), (x + 20, height)], fill="#4A3525")
        draw.ellipse([(x - 30, int(height * 0.2)), (x + 50, int(height * 0.45))], fill="#2E5D32")

    img.save(filepath, quality=88)


def main():
    target_dir = Path("sample_mess").resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    print(f"Creating ~25 messy sample files in: {target_dir}")

    # 1. Screenshots (Pillow)
    create_image_with_ui(
        target_dir / "Screenshot 2026-10-21 at 3.12.png",
        title="Slack #engineering - Deployment Hotfix",
        lines=[
            ("Alex (14:32): Hey team, did the database migration complete on prod?", "#89B4FA"),
            ("Jordan (14:34): Yes, verified table migrations 004_add_user_roles.", "#A6E3A1"),
            ("Sara (14:35): Awesome! Monitoring CPU usage, looking completely green.", "#F9E2AF"),
            ("Logs: [200 OK] GET /api/v1/health - 2.4ms latency", "#CBA6F7")
        ]
    )

    create_image_with_ui(
        target_dir / "Screen Shot 2026-09-14 at 10.45.22 PM.png",
        title="VS Code - Compiler Error TS2322",
        lines=[
            ("Error in src/auth/session.ts(42,15):", "#F38BA8"),
            ("Type 'string | null' is not assignable to type 'string'.", "#F38BA8"),
            ("  41 |   const token = getBearerToken(req);", "#CDD6F4"),
            ("  42 |   validateSession(token);", "#FAB387"),
            ("                       ^^^^^", "#F38BA8"),
            ("Quick Fix: Add null assertion or if check.", "#A6ADC8")
        ]
    )

    create_image_with_ui(
        target_dir / "Screenshot_20261005_dashboard.png",
        title="Analytics Dashboard - Monthly Active Users",
        lines=[
            ("Monthly Active Users: 142,850 (+12.4% MoM)", "#A6E3A1"),
            ("Server Uptime: 99.98% across all 4 regions", "#89B4FA"),
            ("API Latency p95: 42ms | Error Rate: 0.01%", "#F9E2AF"),
            ("Database Cache Hit Ratio: 94.2%", "#CBA6F7")
        ]
    )

    create_image_with_ui(
        target_dir / "capture_receipt_payment_stripe.png",
        title="Stripe - Payment Confirmation #ch_3Ng7Xb",
        lines=[
            ("Payment succeeded: $89.00 USD", "#A6E3A1"),
            ("Customer: user@example.com", "#CDD6F4"),
            ("Description: Pro Plan Subscription (Annual)", "#89B4FA"),
            ("Status: Paid on Oct 18, 2026", "#F9E2AF")
        ]
    )

    # 2. Photos & Graphics
    create_nature_photo(target_dir / "IMG_4821.jpg")
    create_nature_photo(target_dir / "DCIM_0092_nature.jpg")

    # Meme / graphic
    meme_img = Image.new("RGB", (500, 350), color="#FFF3CD")
    d_meme = ImageDraw.Draw(meme_img)
    d_meme.rectangle([(10, 10), (490, 340)], outline="#856404", width=3)
    d_meme.text((30, 40), "WHEN THE CODE COMPILES ON FIRST TRY", fill="#856404")
    d_meme.rectangle([(150, 100), (350, 220)], fill="#FFEBAA")
    d_meme.text((180, 150), "\\(^o^)/", fill="#533F03")
    d_meme.text((120, 260), "*sloth raises hands in celebration*", fill="#856404")
    meme_img.save(target_dir / "meme_cat_coffee_v1.jpg")

    avatar_img = Image.new("RGB", (200, 200), color="#3B82F6")
    d_av = ImageDraw.Draw(avatar_img)
    d_av.ellipse([(40, 40), (160, 160)], fill="#FFFFFF")
    avatar_img.save(target_dir / "profile_photo_avatar.png")

    # 3. Receipts & Bills
    (target_dir / "electric_bill_unpaid_urgent.txt").write_text(
        "PACIFIC ELECTRIC & POWER COMPANY\n"
        "Account Number: 8492-1049-22\n"
        "Due Date: November 05, 2026\n"
        "Total Amount Due: $148.65\n"
        "Billing Period: Sep 20, 2026 - Oct 20, 2026\n"
        "Service Address: 742 Evergreen Terrace\n"
        "Status: UNPAID - Please pay immediately to avoid late fee.\n",
        encoding="utf-8"
    )

    (target_dir / "receipt_target_groceries_oct.txt").write_text(
        "TARGET STORE #1240 - SAN FRANCISCO CA\n"
        "Date: 10/18/2026 18:42\n"
        "Cashier: Maya S.\n\n"
        "Items:\n"
        "1. Organic Oat Milk 64oz        $4.99\n"
        "2. Sourdough Bread Whole        $3.49\n"
        "3. Honeycrisp Apples 3lb        $5.99\n"
        "4. Dark Roast Ground Coffee    $10.49\n"
        "5. Extra Virgin Olive Oil       $12.99\n"
        "Subtotal:                      $37.95\n"
        "Tax (8.5%):                     $3.23\n"
        "Total Paid (Visa *4920):       $41.18\n"
        "Thank you for shopping at Target!\n",
        encoding="utf-8"
    )

    (target_dir / "aws_cloud_invoice_0921.txt").write_text(
        "AMAZON WEB SERVICES INVOICE\n"
        "Invoice Number: INV-92817412\n"
        "Billing Date: October 01, 2026\n\n"
        "Charges Breakdown:\n"
        "- Amazon Elastic Compute Cloud (EC2):  $24.50\n"
        "- Simple Storage Service (S3):          $4.20\n"
        "- Amazon Relational Database (RDS):     $18.00\n"
        "- Route 53 Hosted Zones:                $1.50\n"
        "Total Amount Billed:                   $48.20 USD\n"
        "Payment Status: Paid automatically via Amex *1004\n",
        encoding="utf-8"
    )

    (target_dir / "starbucks_coffee_receipt.txt").write_text(
        "STARBUCKS STORE #0492\n"
        "Order #34 - 10/19/2026 08:15 AM\n"
        "1 Grande Caramel Macchiato     $5.75\n"
        "  - Oat milk substitute        $0.80\n"
        "1 Butter Croissant             $3.45\n"
        "Total:                        $10.00\n"
        "Tip:                           $1.50\n"
        "Grand Total:                  $11.50\n",
        encoding="utf-8"
    )

    # 4. Documents (DOCX)
    doc1 = docx.Document()
    doc1.add_heading("Freelance Design & Development Agreement", level=1)
    doc1.add_paragraph("This Agreement is made on October 15, 2026, by and between Nexus Labs ('Client') and Studio Craft ('Contractor').")
    doc1.add_heading("1. Scope of Work", level=2)
    doc1.add_paragraph("The Contractor agrees to perform web redesign, brand identity overhaul, and responsive frontend implementation.")
    doc1.add_heading("2. Payment Terms", level=2)
    doc1.add_paragraph("Total project fee is $6,500 payable in three milestone installments upon deliverable review.")
    doc1.save(str(target_dir / "final_final_v2.docx"))

    doc2 = docx.Document()
    doc2.add_heading("Curriculum Vitae - Senior Software Engineer", level=1)
    doc2.add_paragraph("Alex Vance | alex.vance@example.com | San Francisco, CA")
    doc2.add_heading("Professional Summary", level=2)
    doc2.add_paragraph("Full-stack engineer with 7+ years of experience specializing in Python, FastAPI, distributed systems, and local LLM orchestration.")
    doc2.add_heading("Experience", level=2)
    doc2.add_paragraph("Senior Engineer at CloudCorp (2022 - Present): Architected microservices serving 10M+ daily requests.")
    doc2.save(str(target_dir / "resume_draft_updated.docx"))

    # 5. Documents & Study Notes (PDF)
    create_pdf(
        target_dir / "lecture_4_quantum_algorithms.pdf",
        title="CS 280: Quantum Algorithms & Complexity",
        paragraphs=[
            "Lecture 4: Shor Factoring and Grover Search Speedups.",
            "In this lecture, we introduce quantum phase estimation and polynomial speedups.",
            "Grover search achieves O(sqrt(N)) query complexity for unstructured databases.",
            "Read Chapter 6 of Nielsen and Chuang for homework assignment due Friday."
        ]
    )

    create_pdf(
        target_dir / "contract_nda_confidentiality.pdf",
        title="Mutual Non-Disclosure Agreement (NDA)",
        paragraphs=[
            "This Mutual Non-Disclosure Agreement is entered into on October 10, 2026.",
            "Both parties agree to hold all Proprietary Information in strict confidence.",
            "Proprietary Information includes technical architecture, algorithms, and models.",
            "The term of confidentiality shall extend 3 years from disclosure."
        ]
    )

    create_pdf(
        target_dir / "quarterly_financial_report_q3.pdf",
        title="Q3 2026 Corporate Financial Review",
        paragraphs=[
            "Executive Summary: Revenue grew 18% year-over-year reaching $4.2M.",
            "Gross margins improved by 250 basis points due to infrastructure efficiencies.",
            "Cash runway remains healthy at 28 months with zero debt outstanding.",
            "Guidance for Q4 anticipates sustained expansion into enterprise accounts."
        ]
    )

    # 6. Study Notes (Text / Markdown)
    (target_dir / "biology_mitosis_notes.txt").write_text(
        "Biology 101 - Cell Division & Mitosis Summary\n"
        "--------------------------------------------\n"
        "1. Interphase: DNA replicates in S phase, centrosomes duplicate.\n"
        "2. Prophase: Chromatin condenses into distinct chromosomes, spindle forms.\n"
        "3. Metaphase: Chromosomes align along the metaphase plate.\n"
        "4. Anaphase: Sister chromatids pull apart toward opposite poles.\n"
        "5. Telophase: Nuclear envelopes reform, cytokinesis divides cytoplasm.\n",
        encoding="utf-8"
    )

    (target_dir / "history_ch7_renaissance.md").write_text(
        "# History Chapter 7: The Italian Renaissance\n\n"
        "## Key Factors:\n"
        "- Rise of wealthy merchant families in Florence (e.g., Medici).\n"
        "- Humanism and rediscovery of classical Greek and Roman texts.\n"
        "- Innovations in perspective, anatomy, and oil painting (Da Vinci, Michelangelo).\n"
        "- The printing press by Gutenberg (c. 1440) drastically accelerated knowledge dissemination.\n",
        encoding="utf-8"
    )

    # 7. Code Scripts
    (target_dir / "temp_test_script.py").write_text(
        "import fastapi\n"
        "import uvicorn\n\n"
        "app = fastapi.FastAPI()\n\n"
        "@app.get('/ping')\n"
        "def ping():\n"
        "    return {'status': 'pong', 'version': '1.0'}\n\n"
        "if __name__ == '__main__':\n"
        "    uvicorn.run(app, host='127.0.0.1', port=9000)\n",
        encoding="utf-8"
    )

    (target_dir / "quick_fix_auth.js").write_text(
        "// Hotfix for JWT signature verification\n"
        "const jwt = require('jsonwebtoken');\n\n"
        "function verifyAuthHeader(req, res, next) {\n"
        "  const auth = req.headers['authorization'];\n"
        "  if (!auth) return res.status(401).json({ error: 'Missing auth' });\n"
        "  const token = auth.split(' ')[1];\n"
        "  jwt.verify(token, process.env.JWT_SECRET, (err, decoded) => {\n"
        "    if (err) return res.status(403).json({ error: 'Invalid token' });\n"
        "    req.user = decoded;\n"
        "    next();\n"
        "  });\n"
        "}\n\n"
        "module.exports = { verifyAuthHeader };\n",
        encoding="utf-8"
    )

    (target_dir / "database_dump_schema.sql").write_text(
        "-- Core Database Schema Export\n"
        "CREATE TABLE users (\n"
        "    id BIGSERIAL PRIMARY KEY,\n"
        "    email VARCHAR(255) UNIQUE NOT NULL,\n"
        "    password_hash VARCHAR(255) NOT NULL,\n"
        "    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()\n"
        ");\n\n"
        "CREATE TABLE files (\n"
        "    id BIGSERIAL PRIMARY KEY,\n"
        "    user_id BIGINT REFERENCES users(id),\n"
        "    original_name VARCHAR(500) NOT NULL,\n"
        "    category VARCHAR(100) NOT NULL,\n"
        "    size_bytes BIGINT NOT NULL\n"
        ");\n",
        encoding="utf-8"
    )

    (target_dir / "deploy_server.sh").write_text(
        "#!/usr/bin/env bash\n"
        "set -euo pipefail\n\n"
        "echo 'Building production container...'\n"
        "docker build -t myapp:latest .\n"
        "echo 'Restarting service...'\n"
        "docker compose down && docker compose up -d\n"
        "echo 'Deployment complete!'\n",
        encoding="utf-8"
    )

    # 8. Archive & Misc
    zip_path = target_dir / "backup_archive_2026.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("database_backup.txt", "Automated SQL dump backup 2026-10-18")
        zf.writestr("env_sample.txt", "APP_ENV=production\nDEBUG=False")

    (target_dir / "random_notes_todo.md").write_text(
        "# Weekly To-Do & Reminders\n\n"
        "- [x] Pick up dry cleaning\n"
        "- [x] Dentist appointment on Thursday 3 PM\n"
        "- [ ] Buy fresh coffee beans\n"
        "- [ ] Refactor folder organization script\n"
        "- [ ] Call grandma on the weekend\n",
        encoding="utf-8"
    )

    (target_dir / "meeting_minutes_sept.txt").write_text(
        "Team Sync Meeting Minutes - Sep 28, 2026\n"
        "Attendees: Alex, Maya, Carlos, Wei\n\n"
        "Discussion Points:\n"
        "1. Q4 roadmap alignment and milestone deadlines.\n"
        "2. Moving heavy compute tasks to local on-device models to preserve privacy.\n"
        "3. Sloth Sorter prototype feedback: users love the pacing and undo capability.\n",
        encoding="utf-8"
    )

    created_count = len(list(target_dir.glob("*")))
    print(f"Successfully generated {created_count} sample messy files in {target_dir}")


if __name__ == "__main__":
    main()
